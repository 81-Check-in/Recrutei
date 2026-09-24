"""Extração de texto: PDF, DOCX, imagens (OCR) e Google Docs."""
import io
import re
import signal
import threading
import zipfile
from contextlib import contextmanager
from urllib.parse import unquote
from typing import Dict, Optional, Tuple

import requests

from config import (
    log, TAMANHO_MAXIMO_ANEXO, LIMITE_PIXELS_IMAGEM, LIMITE_ZIP_DESCOMPRIMIDO, LIMITE_ZIP_ENTRADAS,
    LADO_MAX_PDF_PX, TEMPO_MAX_OCR, TEMPO_MAX_EXTRACAO,
)
from utils import limpar_texto


class _TempoExcedido(BaseException):
    """BaseException para não ser engolida pelos 'except Exception' internos."""


@contextmanager
def _limite_de_tempo(segundos: int):
    """Interrompe extrações que travam. Só age em Linux/macOS (SIGALRM)."""
    if (not hasattr(signal, "SIGALRM")
            or threading.current_thread() is not threading.main_thread()):
        yield
        return

    def _estourou(signum, frame):
        raise _TempoExcedido()

    anterior = signal.signal(signal.SIGALRM, _estourou)
    signal.alarm(segundos)
    try:
        yield
    finally:
        signal.alarm(0)
        signal.signal(signal.SIGALRM, anterior)


def _de_pdf(conteudo: bytes) -> Tuple[Optional[str], bool]:
    """Retorna (texto, ocr_aplicado)."""
    texto = ""
    pag1_tem_imagem = False

    # 1ª tentativa: extração direta
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(conteudo)) as pdf:
            paginas = [p.extract_text() or "" for p in pdf.pages[:15]]
            texto = "\n".join(paginas)
            if pdf.pages:
                pag1_tem_imagem = bool(pdf.pages[0].images)
    except Exception as e:
        log.debug(f"  pdfplumber falhou: {e}")

    # 2ª tentativa: pypdf
    if len(texto.strip()) < 100:
        try:
            from pypdf import PdfReader
            leitor = PdfReader(io.BytesIO(conteudo))
            texto = "\n".join((p.extract_text() or "") for p in leitor.pages[:15])
        except Exception as e:
            log.debug(f"  pypdf falhou: {e}")

    # 3ª tentativa: OCR (PDF escaneado)
    if len(texto.strip()) < 100:
        log.info("  PDF sem texto — aplicando OCR")
        texto_ocr = _ocr_pdf(conteudo, ultima_pagina=5)
        if texto_ocr and len(texto_ocr.strip()) >= 100:
            return texto_ocr, True
        return (texto or None), False

    # Achou texto, mas a 1ª página tem imagem: modelos de currículo (Canva e afins)
    # costumam "achatar" o cabeçalho (nome, contato, foto) como gráfico — texto
    # normal não pega isso. Soma o OCR só da 1ª página ao invés de descartar o resto.
    if pag1_tem_imagem:
        cabecalho = _ocr_pdf(conteudo, ultima_pagina=1)
        if cabecalho and cabecalho.strip():
            texto = cabecalho.strip() + "\n" + texto

    return texto, False


def _ocr_pdf(conteudo: bytes, ultima_pagina: int = 5) -> Optional[str]:
    try:
        from pdf2image import convert_from_bytes
        import pytesseract
        # size limita o lado maior da página: páginas gigantes não estouram a memória
        imagens = convert_from_bytes(conteudo, size=LADO_MAX_PDF_PX,
                                     first_page=1, last_page=ultima_pagina,
                                     timeout=TEMPO_MAX_OCR)
        partes = [pytesseract.image_to_string(img, lang="por+eng",
                                              timeout=TEMPO_MAX_OCR)
                  for img in imagens]
        return "\n".join(partes)
    except Exception as e:
        log.warning(f"  OCR indisponível ou falhou: {e}")
        return None


def _de_imagem(conteudo: bytes) -> Optional[str]:
    try:
        import pytesseract
        from PIL import Image
        img = Image.open(io.BytesIO(conteudo))
        if img.width * img.height > LIMITE_PIXELS_IMAGEM:
            log.warning("  Imagem grande demais — ignorada")
            return None
        return pytesseract.image_to_string(img, lang="por+eng",
                                           timeout=TEMPO_MAX_OCR)
    except Exception as e:
        log.warning(f"  OCR de imagem falhou: {e}")
        return None


def _zip_seguro(conteudo: bytes) -> bool:
    """Recusa DOCX (zip) que se expandem além do razoável (bomba de zip)."""
    try:
        with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
            itens = z.infolist()
            if len(itens) > LIMITE_ZIP_ENTRADAS:
                return False
            return sum(i.file_size for i in itens) <= LIMITE_ZIP_DESCOMPRIMIDO
    except zipfile.BadZipFile:
        return False


def _de_docx(conteudo: bytes) -> Optional[str]:
    if not _zip_seguro(conteudo):
        log.warning("  DOCX inválido ou grande demais depois de descomprimido")
        return None
    try:
        from docx import Document
        doc = Document(io.BytesIO(conteudo))
        partes = [p.text for p in doc.paragraphs if p.text.strip()]
        # Tabelas costumam conter dados de contato
        for tabela in doc.tables:
            for linha in tabela.rows:
                celulas = [c.text.strip() for c in linha.cells if c.text.strip()]
                if celulas:
                    partes.append(" | ".join(celulas))
        return "\n".join(partes)
    except Exception as e:
        log.warning(f"  Falha ao ler DOCX: {e}")
        return None


def _de_google_docs(url: str) -> Tuple[Optional[str], bool]:
    """
    Baixa o documento se o compartilhamento for público. Retorna (texto, ocr_aplicado).
    O link mais comum (app do Drive no celular) é de um PDF/DOCX enviado ao Drive, não
    um Google Docs nativo: o download devolve os bytes do arquivo, não texto puro, e
    precisa passar pelos mesmos extratores usados para anexo (senão vira texto ilegível
    e a IA rejeita como "não é currículo").
    """
    m = re.search(r"/d/([a-zA-Z0-9_-]+)", url)
    if not m:
        return None, False
    doc_id = m.group(1)

    endpoints = [
        f"https://docs.google.com/document/d/{doc_id}/export?format=txt",
        f"https://drive.google.com/uc?export=download&id={doc_id}",
    ]
    for ep in endpoints:
        try:
            r = requests.get(ep, timeout=20, allow_redirects=True)
            if r.status_code != 200 or len(r.content) < 100:
                continue
            if "accounts.google.com" in r.url:
                continue  # documento privado, pede login

            tipo = r.headers.get("Content-Type", "")
            if "application/pdf" in tipo or r.content[:5] == b"%PDF-":
                texto, ocr = _de_pdf(r.content)
                if texto and len(texto.strip()) >= 100:
                    return texto, ocr
                continue
            if "wordprocessingml" in tipo or r.content[:2] == b"PK":
                texto = _de_docx(r.content)
                if texto and len(texto.strip()) >= 100:
                    return texto, False
                continue
            if "text/html" in tipo:
                continue  # página de aviso/confirmação do Drive, não é o arquivo

            texto = r.content.decode("utf-8", errors="ignore")
            if len(texto.strip()) >= 100:
                return texto, False
        except Exception as e:
            log.debug(f"  Google Docs/Drive ({ep}): {e}")
    return None, False


_MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _e_docx(conteudo: bytes) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
            return "word/document.xml" in z.namelist()
    except Exception:
        return False


def _nome_do_download(resposta, extensao: str) -> str:
    """
    O nome que o Drive informa (Content-Disposition), sem caminho e sempre terminando na extensão do arquivo; sem ele, um nome
    genérico. Prefere a forma filename*=UTF-8''… (percent-encoded); a forma simples vem em UTF-8 mas o requests a lê como
    latin-1 ("currÃ­culo"), então é desfeita.
    """
    cab = resposta.headers.get("Content-Disposition", "")
    nome = ""
    m = re.search(r"filename\*=(?:UTF-8|utf-8)''([^;]+)", cab)
    if m:
        nome = unquote(m.group(1))
    else:
        m = re.search(r'filename="?([^";]+)"?', cab)
        if m:
            nome = m.group(1)
            try:
                nome = nome.encode("latin-1").decode("utf-8")
            except (UnicodeEncodeError, UnicodeDecodeError):
                pass
    nome = re.sub(r"[\\/\x00-\x1f]", "", nome).strip()[:120]
    if not nome:
        return f"curriculo-google-docs{extensao}"
    return nome if nome.lower().endswith(extensao) else nome + extensao


def arquivo_do_google_docs(url: str) -> Optional[Dict]:
    """
    O ARQUIVO original de um link do Google Docs/Drive, para ser guardado e o RH poder abri-lo no painel (só o texto não
    basta: sem o arquivo o painel diz "arquivo original não disponível"). Google Docs nativo: exportado em PDF. Arquivo
    enviado ao Drive: os próprios bytes do PDF ou DOCX. Devolve {"conteudo", "tipo_mime", "nome", "tamanho"}, ou None
    (documento privado, exportação bloqueada, formato que não é PDF nem DOCX, pequeno ou grande demais). Nunca levanta erro:
    sem o arquivo o currículo entra igual, só sem o botão de abrir.
    """
    m = re.search(r"/d/([a-zA-Z0-9_-]+)", url or "")
    if not m:
        return None
    doc_id = m.group(1)
    for ep in (f"https://docs.google.com/document/d/{doc_id}/export?format=pdf",
               f"https://drive.google.com/uc?export=download&id={doc_id}"):
        try:
            r = requests.get(ep, timeout=30, allow_redirects=True)
            if r.status_code != 200 or "accounts.google.com" in r.url:
                continue                                  # não existe como Google Docs, ou é privado (pede login)
            conteudo = r.content
            if not 100 <= len(conteudo) <= TAMANHO_MAXIMO_ANEXO:
                continue
            if conteudo[:5] == b"%PDF-":
                tipo, ext = "application/pdf", ".pdf"
            elif conteudo[:2] == b"PK" and _e_docx(conteudo):
                tipo, ext = _MIME_DOCX, ".docx"
            else:
                continue                                  # página de aviso do Drive (HTML), texto puro etc.
            return {"conteudo": conteudo, "tipo_mime": tipo, "nome": _nome_do_download(r, ext), "tamanho": len(conteudo)}
        except Exception as e:
            log.debug(f"  Arquivo do Google Docs/Drive ({ep}): {type(e).__name__}")
    return None


def _extrair(conteudo: bytes, tipo_mime: str) -> Tuple[Optional[str], bool]:
    if tipo_mime == "application/pdf":
        return _de_pdf(conteudo)
    if "wordprocessingml" in tipo_mime:
        return _de_docx(conteudo), False
    if tipo_mime == "application/msword":
        # .doc antigo: tenta como docx, senão OCR não se aplica
        texto = _de_docx(conteudo)
        return texto, False
    if tipo_mime.startswith("image/"):
        return _de_imagem(conteudo), True
    return None, False


def extrair(conteudo: bytes, tipo_mime: str) -> Tuple[Optional[str], bool]:
    """Roteia para o extrator adequado. Retorna (texto, ocr_aplicado)."""
    try:
        with _limite_de_tempo(TEMPO_MAX_EXTRACAO):
            return _extrair(conteudo, tipo_mime)
    except _TempoExcedido:
        log.warning(f"  Extração interrompida: passou de {TEMPO_MAX_EXTRACAO}s")
        return None, False
    except MemoryError:
        log.warning("  Extração interrompida: memória insuficiente")
        return None, False


def extrair_google_docs(url: str) -> Tuple[Optional[str], bool]:
    """Retorna (texto, ocr_aplicado)."""
    texto, ocr = _de_google_docs(url)
    return (limpar_texto(texto) if texto else None), ocr
